import json, pytest
from conftest import OP, AP, AP2, ADMIN, make
from core.support import GateError, PermissionDenied, Actor, verify_audit_chain
from core.models import AuditEntry, ValidationResult
from core.reports import build_report, export_json, export_xlsx
from core.assessment import evaluate_compatibility, score_readiness


# ---------- readiness / assessment
def test_clean_scenario_is_green_and_allows_plan(wb):
    pid = make(wb)
    r = wb.latest_readiness(pid)
    assert r.status == "GREEN" and r.score == 100.0 and not r.blockers
    assert wb.gate(pid)["allowed"]


def test_issues_scenario_is_red_and_blocks_plan(wb):
    pid = make(wb, "issues")
    r = wb.latest_readiness(pid)
    assert r.status == "RED"
    assert any("no_blocking_artifacts" in b for b in r.blockers)
    assert any("no_blocking_integrations" in b for b in r.blockers)
    with pytest.raises(GateError, match="readiness gate"):
        wb.generate_plan(OP, pid)


def test_compatibility_classification(wb):
    pid = make(wb, "issues")
    f = {(x["artifact"], x["detail"]): x["status"] for x in wb.latest_compat(pid).findings}
    assert f[("custom_java_integration", "LegacyERPBridge")] == "redesign_required"
    assert f[("legacy_report_format", "OldBookPack")] == "unsupported"
    assert f[("file_based_integration", "Treasury_FileDrop")] == "remap_required"
    assert f[("forms", "85 item(s)")] == "supported"


def test_approved_remediation_clears_blocks_and_amber_needs_acceptance(wb):
    # remediation approved -> no critical blockers; unowned integration (-5), unreviewed privileged access (-5),
    # undefined data scope (-7.5) and no cutover attestations (-5) leave 77.5 -> AMBER
    pid = make(wb, "issues", remediation_approved=True, privileged_access_reviewed=False,
               data_scope="", backup_plan=False, rollback_plan="", cutover_window="")
    r = wb.latest_readiness(pid)
    assert not r.blockers and r.score == 77.5 and r.status == "AMBER"
    assert not wb.gate(pid)["allowed"]
    with pytest.raises(PermissionDenied): wb.approve(OP, pid, "readiness_amber")
    wb.approve(AP, pid, "readiness_amber", comment="accepted")
    assert wb.gate(pid)["allowed"]
    wb.generate_plan(OP, pid)


def test_unknown_artifact_is_review_required():
    cfg = {"default": {"status": "review_required", "action": "x", "group": "artifact"}, "rules": []}
    out = evaluate_compatibility({"counts": {"mystery": 3}, "custom_items": []}, cfg)
    assert out[0]["status"] == "review_required"


def test_connectivity_failure_blocks(wb):
    pid = make(wb, tgt_faults={"auth_fail": True})
    r = wb.latest_readiness(pid)
    assert r.status == "RED" and any("connectivity" in b for b in r.blockers)


def test_unreachable_target_records_fail_and_blocks(wb):
    pid = make(wb, tgt_faults={"unreachable": True})
    assert any(c.status == "FAIL" for c in wb.connectivity_latest(pid) if c.side == "target")
    assert wb.latest_readiness(pid).status == "RED"


def test_target_older_version_is_critical(wb):
    pid = make(wb)
    with wb.session() as s:
        p = wb._project(s, pid); p.target = {**p.target, "version": "25.01"}
    import shutil, pathlib
    shutil.rmtree(pathlib.Path(wb.state_dir) / f"p{pid}_target")
    wb.collect_inventory(OP, pid); wb.assess(OP, pid)
    assert any("target_version_compatible" in b for b in wb.latest_readiness(pid).blockers)


def test_capacity_exceeded_is_critical(wb):
    pid = make(wb)
    with wb.session() as s:
        from core.models import Inventory
        inv = wb._latest_inv(s, pid, "target"); inv["capacity"] = {"max_intersections": 1000}
        s.add(Inventory(project_id=pid, side="target", data=inv))
    wb.assess(OP, pid)
    assert any("data_volume_within_capacity" in b for b in wb.latest_readiness(pid).blockers)


def test_missing_cutover_attestations_lower_score_without_blocking(wb):
    pid = make(wb, rollback_plan="", cutover_window="")
    r = wb.latest_readiness(pid)
    assert not r.blockers and r.status == "GREEN" and r.score == 96.7   # 2 of 3 cutover checks fail: -3.3
    assert r.categories["cutover"]["score"] == 33.3


# ---------- plan + execution
def full_run(wb, pid):
    wb.generate_plan(OP, pid)
    res = wb.run_all(OP, pid)
    return res


def test_happy_path_end_to_end(wb):
    pid = make(wb)
    res = full_run(wb, pid)
    assert [k for k, _ in res["ran"]] == ["backup_source", "validate_target", "export_snapshot", "transfer_snapshot"]
    assert "Approval required for step 'import_snapshot'" in res["stopped"]
    wb.approve(AP, pid, "step:import_snapshot")
    res = wb.run_all(OP, pid)
    assert [k for k, _ in res["ran"]] == ["import_snapshot", "technical_validation", "data_reconciliation"]
    assert "business_signoff" in res["stopped"]
    wb.approve(AP2, pid, "step:business_signoff", comment="signed")
    res = wb.run_all(OP, pid)
    assert res["ran"] == [("business_signoff", "COMPLETE")]
    assert all(s.state == "COMPLETE" for s in wb.steps(pid))
    assert wb.get_project(pid).status == "COMPLETED"
    val = wb.rows(ValidationResult, pid)
    assert val and all(v.status == "PASS" for v in val)
    assert {v.category for v in val} == {"technical", "business"}


def test_segregation_of_duties(wb):
    pid = make(wb)
    full_run(wb, pid)
    wb.approve(ADMIN, pid, "step:import_snapshot")
    with pytest.raises(GateError, match="Segregation"):
        wb.run_next(ADMIN, pid)
    wb.approve(AP, pid, "step:import_snapshot")  # a different approver unblocks
    assert wb.run_next(ADMIN, pid).state == "COMPLETE"


def test_rbac_enforced(wb):
    pid = make(wb)
    with pytest.raises(PermissionDenied): wb.run_next(AP, pid)
    with pytest.raises(PermissionDenied): wb.generate_plan(AP, pid)
    with pytest.raises(PermissionDenied): wb.create_project(AP, "x", {}, {})


def test_cannot_execute_without_plan_or_when_paused(wb):
    pid = make(wb)
    with pytest.raises(GateError): wb.run_next(OP, pid)
    wb.generate_plan(OP, pid)
    wb.set_paused(OP, pid, True)
    with pytest.raises(GateError, match="paused"): wb.run_next(OP, pid)
    wb.set_paused(OP, pid, False)
    assert wb.run_next(OP, pid).key == "backup_source"


def test_transient_failure_then_retry_succeeds(wb):
    pid = make(wb, tgt_faults={"fail_once": ["upload_snapshot"]})
    wb.generate_plan(OP, pid)
    res = wb.run_all(OP, pid)
    assert res["ran"][-1] == ("transfer_snapshot", "FAILED")
    st = {s.key: s for s in wb.steps(pid)}["transfer_snapshot"]
    assert st.failure_class == "TIMEOUT" and st.attempts == 1
    with pytest.raises(GateError, match="retry"): wb.run_next(OP, pid)
    wb.retry_step(OP, pid, "transfer_snapshot")
    assert wb.run_next(OP, pid).state == "COMPLETE"


def test_permanent_failure_exhausts_retries(wb):
    pid = make(wb, src_faults={"fail_always": ["export_snapshot"]})
    wb.generate_plan(OP, pid)
    wb.run_all(OP, pid)  # backup_source also uses export_snapshot -> fails first
    for _ in range(2):
        wb.retry_step(OP, pid, "backup_source"); wb.run_next(OP, pid)
    with pytest.raises(GateError, match="not retryable"): wb.retry_step(OP, pid, "backup_source")


def test_validation_failure_detects_dropped_artifacts_and_rerun_fixes_nothing_until_cause_fixed(wb):
    pid = make(wb, tgt_faults={"drop_on_import": {"forms": 3}})
    full_run(wb, pid); wb.approve(AP, pid, "step:import_snapshot")
    res = wb.run_all(OP, pid)
    assert res["ran"][-1] == ("technical_validation", "FAILED")
    st = {s.key: s for s in wb.steps(pid)}["technical_validation"]
    assert st.failure_class == "VALIDATION" and "forms" in st.message
    bad = [v for v in wb.rows(ValidationResult, pid) if v.status == "FAIL"]
    assert [(v.check, v.diff) for v in bad] == [("forms", -3)]
    # nothing downstream may run, sign-off included
    assert {s.key: s.state for s in wb.steps(pid)}["business_signoff"] == "BLOCKED"


def test_control_total_mismatch_fails_reconciliation_and_rerun_requires_new_approval(wb):
    pid = make(wb, tgt_faults={"skew_control_total": {"Revenue": 500.0}})
    full_run(wb, pid); wb.approve(AP, pid, "step:import_snapshot")
    res = wb.run_all(OP, pid)
    assert res["ran"][-1] == ("data_reconciliation", "FAILED")
    wb.rerun_from(OP, pid, "import_snapshot")
    states = {s.key: s.state for s in wb.steps(pid)}
    assert states["import_snapshot"] == "AWAITING_APPROVAL" and states["data_reconciliation"] == "BLOCKED"
    with pytest.raises(GateError, match="Approval required"): wb.run_next(OP, pid)


def test_skip_refused_for_unsafe_step(wb):
    pid = make(wb); wb.generate_plan(OP, pid)
    with pytest.raises(GateError, match="not safe to skip"): wb.skip_step(OP, pid, "import_snapshot", "x")


def test_plan_cannot_be_regenerated_after_execution_started(wb):
    pid = make(wb); wb.generate_plan(OP, pid); wb.run_next(OP, pid)
    with pytest.raises(GateError, match="already executing"): wb.generate_plan(OP, pid)


def test_gate_is_rechecked_at_execution_time(wb):
    pid = make(wb); wb.generate_plan(OP, pid); wb.run_next(OP, pid)
    wb.update_settings(OP, pid, remediation_approved=False)
    with wb.session() as s:
        p = wb._project(s, pid); p.target = {**p.target, "faults": {"auth_fail": True}}
    wb.run_connectivity(OP, pid); wb.assess(OP, pid)   # new evidence turns the gate RED mid-migration
    with pytest.raises(GateError, match="blocks execution"): wb.run_next(OP, pid)


def test_approval_subject_validation(wb):
    pid = make(wb); wb.generate_plan(OP, pid)
    with pytest.raises(GateError): wb.approve(AP, pid, "step:backup_source")
    with pytest.raises(GateError): wb.approve(AP, pid, "bogus")


# ---------- audit + evidence + reports
def test_audit_chain_detects_tampering(wb):
    pid = make(wb); full_run(wb, pid)
    with wb.session() as s:
        assert verify_audit_chain(s) == (True, None)
        e = s.query(AuditEntry).filter(AuditEntry.action == "step.run").first()
        e.details = {**e.details, "state": "COMPLETE-FORGED"}
    with wb.session() as s:
        ok, bad = verify_audit_chain(s)
        assert not ok and bad is not None


def test_report_exports(wb, tmp_path):
    pid = make(wb); full_run(wb, pid)
    wb.approve(AP, pid, "step:import_snapshot"); wb.run_all(OP, pid)
    r = build_report(wb, pid)
    assert r["audit_chain_intact"] and r["readiness"]["status"] == "GREEN"
    assert len(r["evidence"]) >= 7 and all(len(e["sha256"]) == 64 for e in r["evidence"])
    j = json.loads(open(export_json(wb, pid, tmp_path / "r.json")).read())
    assert j["project"]["id"] == pid
    from openpyxl import load_workbook
    book = load_workbook(export_xlsx(wb, pid, tmp_path / "r.xlsx"))
    assert {"Summary", "Readiness", "Compatibility", "Plan & Execution", "Validation", "Approvals", "Evidence", "Audit"} <= set(book.sheetnames)
    assert book["Plan & Execution"].max_row == 9


def test_no_secrets_stored(wb):
    pid = make(wb)
    blob = json.dumps([wb.get_project(pid).source, wb.get_project(pid).target]).lower()
    assert "password" not in blob
