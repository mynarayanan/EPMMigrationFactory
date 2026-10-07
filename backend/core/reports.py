"""Assessment / execution / validation evidence report (JSON + Excel)."""
from __future__ import annotations
import json
from pathlib import Path
from sqlalchemy import select
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .models import (Approval, Evidence, ValidationResult, AuditEntry, RunLog)
from .support import verify_audit_chain


def build_report(wb, pid: int) -> dict:
    p = wb.get_project(pid)
    rr, cr = wb.latest_readiness(pid), wb.latest_compat(pid)
    with wb.session() as s:
        ok, bad = verify_audit_chain(s)
        audit_rows = [e for e in s.scalars(select(AuditEntry).where(AuditEntry.project_id == pid).order_by(AuditEntry.id))]
    return {
        "project": {"id": p.id, "name": p.name, "client": p.client, "pattern": p.pattern, "status": p.status,
                    "owner": p.owner, "planned_date": p.planned_date, "created_by": p.created_by},
        "readiness": None if not rr else {"score": rr.score, "status": rr.status, "blockers": rr.blockers,
                                          "categories": rr.categories, "assessed_at": str(rr.created_at)},
        "compatibility": [] if not cr else cr.findings,
        "gate": wb.gate(pid),
        "steps": [{"seq": x.seq, "key": x.key, "name": x.name, "state": x.state, "attempts": x.attempts,
                   "failure_class": x.failure_class, "message": x.message,
                   "started_at": str(x.started_at or ""), "finished_at": str(x.finished_at or "")} for x in wb.steps(pid)],
        "validation": [{"category": v.category, "check": v.check, "source": v.source_value, "target": v.target_value,
                        "diff": v.diff, "status": v.status} for v in wb.rows(ValidationResult, pid)],
        "approvals": [{"subject": a.subject, "approver": a.approver, "decision": a.decision, "comment": a.comment,
                       "at": str(a.created_at)} for a in wb.rows(Approval, pid)],
        "evidence": [{"step": e.step_key, "kind": e.kind, "sha256": e.sha256, "at": str(e.ts)} for e in wb.rows(Evidence, pid)],
        "audit": [{"id": e.id, "ts": str(e.ts), "actor": e.actor, "action": e.action, "details": e.details} for e in audit_rows],
        "audit_chain_intact": ok, "audit_chain_first_bad_id": bad,
    }


def export_json(wb, pid, path) -> str:
    Path(path).write_text(json.dumps(build_report(wb, pid), indent=2, default=str)); return str(path)


def export_xlsx(wb, pid, path) -> str:
    r = build_report(wb, pid)
    book = Workbook()
    head = PatternFill("solid", fgColor="1F3A5F")

    def sheet(name, rows, cols, first=False):
        ws = book.active if first else book.create_sheet()
        ws.title = name
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for row in rows: ws.append([str(row.get(c, "")) if isinstance(row.get(c), (dict, list)) else row.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(c) + 4)
        return ws

    pr = r["project"]; rd = r["readiness"] or {}
    sheet("Summary", [{"item": k, "value": v} for k, v in {**pr, "readiness_score": rd.get("score"),
          "readiness_status": rd.get("status"), "gate_allowed": r["gate"]["allowed"],
          "audit_chain_intact": r["audit_chain_intact"]}.items()], ["item", "value"], first=True)
    sheet("Readiness", [{"category": c, "weight": v["weight"], "score": v["score"], "check": k["id"],
                         "passed": k["passed"], "critical": k["critical"], "detail": k["detail"]}
                        for c, v in (rd.get("categories") or {}).items() for k in v["checks"]],
          ["category", "weight", "score", "check", "passed", "critical", "detail"])
    sheet("Compatibility", r["compatibility"], ["artifact", "detail", "status", "action", "group"])
    sheet("Plan & Execution", r["steps"], ["seq", "key", "name", "state", "attempts", "failure_class", "message", "started_at", "finished_at"])
    sheet("Validation", r["validation"], ["category", "check", "source", "target", "diff", "status"])
    sheet("Approvals", r["approvals"], ["subject", "approver", "decision", "comment", "at"])
    sheet("Evidence", r["evidence"], ["step", "kind", "sha256", "at"])
    sheet("Audit", r["audit"], ["id", "ts", "actor", "action", "details"])
    book.save(path); return str(path)
